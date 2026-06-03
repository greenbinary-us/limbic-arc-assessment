import re
from pptx import Presentation
from pptx.util import Pt

PPTX = r'C:\Users\vmank\Documents\Code\GBClients\LimbicArc-assessment\limbic-arc-assessment\presentation\Limbic Arc Tech Assessment -  Briefing.pptx'

prs = Presentation(PPTX)

# Match " debt" or "Debt" as a whole word but NOT when already preceded by "tech "
pattern = re.compile(r'(?<!tech )\bdebt\b', re.IGNORECASE)

def fix(text):
    def repl(m):
        # preserve original case of 'D'/'d'
        return 'tech ' + m.group(0)
    return pattern.sub(repl, text)

changes = 0

def process_tf(tf):
    global changes
    for para in tf.paragraphs:
        for run in para.runs:
            original = run.text
            updated = fix(original)
            if updated != original:
                run.text = updated
                changes += 1
                print(f"  [{repr(original)}] -> [{repr(updated)}]")

for i, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.has_text_frame:
            process_tf(shape.text_frame)
        if shape.shape_type == 19:  # table
            for row in shape.table.rows:
                for cell in row.cells:
                    process_tf(cell.text_frame)
    # speaker notes
    try:
        notes_tf = slide.notes_slide.notes_text_frame
        process_tf(notes_tf)
    except Exception:
        pass

print(f"\nTotal replacements: {changes}")
prs.save(PPTX)
print("Saved.")
