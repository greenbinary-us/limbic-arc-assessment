import os, struct
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
IMG  = os.path.join(HERE, "..", "images")
OUT  = os.path.join(HERE, "..", "Limbic Arc - CEO Briefing (v0.1).pptx")

NAVY  = RGBColor(0x1F,0x3A,0x5F)
NAVY2 = RGBColor(0x2C,0x52,0x7E)
INK   = RGBColor(0x16,0x20,0x2B)
SLATE = RGBColor(0x5C,0x6B,0x7A)
MUTED = RGBColor(0x8A,0x97,0xA4)
TEAL  = RGBColor(0x1A,0xA1,0x92)
RISK  = RGBColor(0xD9,0x53,0x4F)
AMBER = RGBColor(0xE8,0xA3,0x3D)
EXIGO = RGBColor(0x2E,0x7D,0x5B)
FLUID = RGBColor(0x6C,0x5C,0xE7)
LIGHT = RGBColor(0xF3,0xF6,0xF9)
PANEL = RGBColor(0xEA,0xF0,0xF6)
WHITE = RGBColor(0xFF,0xFF,0xFF)
LINE  = RGBColor(0xD7,0xDE,0xE6)
FONT  = "Segoe UI"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, l, t, w, h, fill=None, line=None, line_w=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w if line_w else 1)
    return sp

def para(tf, segs, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
         space_after=8, space_before=0, line=1.12, first=False, bullet=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    if isinstance(segs, str):
        segs = [(segs, bold, color)]
    for seg in segs:
        txt = seg[0]
        b = seg[1] if len(seg) > 1 else bold
        c = seg[2] if len(seg) > 2 else color
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = c; r.font.name = FONT
    return p

def textbox(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf, m, Pt(2))
    return tf

def chrome(s, n, title="Limbic Arc — Platform Assessment · CEO Briefing"):
    tf = textbox(s, 0.7, 7.06, 9, 0.35)
    para(tf, title, 9.5, MUTED, first=True)
    tf2 = textbox(s, 11.8, 7.06, 1.0, 0.35)
    para(tf2, str(n), 9.5, MUTED, align=PP_ALIGN.RIGHT, first=True)

def title_block(s, title, accent=TEAL):
    tf = textbox(s, 0.7, 0.5, 12, 1.0)
    para(tf, title, 31, NAVY, bold=True, first=True, line=1.05)
    rect(s, 0.72, 1.42, 1.7, 0.07, fill=accent)

def png_size(path):
    with open(path, "rb") as f:
        d = f.read(33)
    return struct.unpack(">II", d[16:24])

def image_slide(n, fname):
    s = slide()
    rect(s, 0, 0, SW, SH, fill=WHITE)
    path = os.path.join(IMG, fname)
    iw, ih = png_size(path)
    ar = iw/ih
    bl, bt, bw, bh = 0.35, 0.32, 12.63, 6.55
    if ar > bw/bh:
        w = bw; h = bw/ar
    else:
        h = bh; w = bh*ar
    l = bl + (bw-w)/2; t = bt + (bh-h)/2
    s.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))
    return s

# ---------------------------------------------------------------- SLIDE 1
s = slide()
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, 0, 0.28, SH, fill=TEAL)
tf = textbox(s, 1.1, 2.35, 11, 3.0)
para(tf, "Limbic Arc", 54, WHITE, bold=True, first=True, space_after=4)
para(tf, "Platform Assessment", 30, RGBColor(0xCF,0xD9,0xE6), bold=False, space_after=2)
para(tf, [("A plain-language briefing for leadership", False, TEAL)], 20, space_after=0)
tf2 = textbox(s, 1.12, 6.15, 11, 0.8)
para(tf2, [("Companion to the detailed Technical Assessment (v0.4)", False, RGBColor(0x9A,0xAB,0xBD))], 13, first=True)

# ---------------------------------------------------------------- SLIDE 2  bottom line
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "The bottom line")
cards = [
    ("Two good platforms — neither is the problem.",
     "Exigo runs your distributor network, ranks and commissions. Fluid is your modern online store. Keep both.", NAVY),
    ("The problem is the gap between them.",
     "The seam between the two systems, the aging software your websites are built on, and having no in-house team to change any of it.", RISK),
    ("Recommendation: modernize gradually and safely.",
     "One piece at a time. The first step is small, low-risk, and starts paying off in weeks — but it needs a team in place.", TEAL),
]
y = 1.85
for head, body, col in cards:
    rect(s, 0.7, y, 0.12, 1.45, fill=col)
    tf = textbox(s, 1.0, y, 11.4, 1.45, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(head, True, INK)], 21, first=True, space_after=4)
    para(tf, [(body, False, SLATE)], 16, line=1.2)
    y += 1.72

# ---------------------------------------------------------------- SLIDE 3  scope
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "What we looked at")
tf = textbox(s, 0.75, 1.9, 11.8, 4.6)
para(tf, [("We assessed the surfaces your people actually touch — plus the plumbing between your two platforms:", False, SLATE)], 17, first=True, space_after=14, line=1.2)
for a,b in [
    ("Your customer-facing site, ", "customer account area, and distributor (affiliate) back office"),
    ("The connection ", "between Exigo and Fluid"),
    ("The concerns leadership cares about: ", "billing reliability, security, customer experience, speed of change, and insights"),
]:
    para(tf, [("•  ", True, TEAL), (a, True, INK), (b, False, INK)], 19, space_after=11, line=1.2)
rect(s, 0.75, 5.85, 11.8, 0.9, fill=LIGHT)
tf2 = textbox(s, 1.0, 5.85, 11.3, 0.9, anchor=MSO_ANCHOR.MIDDLE)
para(tf2, [("Out of scope:  ", True, NAVY), ("your core Web App — except that customers shouldn't have to ", False, SLATE), ("log in repeatedly", True, INK), (" to reach it.", False, SLATE)], 15.5, first=True, line=1.2)

# ---------------------------------------------------------------- SLIDE 4  current state
image_slide(4, "current-state.png")

# ---------------------------------------------------------------- SLIDE 5  five things (table)
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "The five things that matter")
rows = [
    ("#", "What's happening", "Why it matters to the business"),
    ("1", "Payment & subscription data is split between Exigo and Fluid; no single “card on file.”",
          "Direct risk to revenue & trust — fragile subscriptions, billing inconsistencies, extra compliance cost.  #1 to fix."),
    ("2", "You depend on your vendor to change anything (old, tightly-wound website software).",
          "You don't control your own speed or cost. Your roadmap runs on someone else's calendar."),
    ("3", "Three different website experiences and repeated logins.",
          "Brand & trust erosion, lower conversion, more support tickets."),
    ("4", "Missing basics: no proper access control, no single sign-on, analytics that don't drive decisions.",
          "Security & oversight gaps; decisions made without a clear view."),
    ("5", "No in-house technology team.",
          "The root cause of #2. Nothing improves durably until this is addressed."),
]
tbl_shape = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.7), Inches(11.93), Inches(5.0))
tbl = tbl_shape.table
tbl.first_row = False; tbl.horz_banding = False
tbl.columns[0].width = Inches(0.6)
tbl.columns[1].width = Inches(5.5)
tbl.columns[2].width = Inches(5.83)
row_h = [0.55] + [0.9]*5
for i,h in enumerate(row_h):
    tbl.rows[i].height = Inches(h)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Pt(8); cell.margin_right = Pt(8)
        cell.margin_top = Pt(5); cell.margin_bottom = Pt(5)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if ci==0 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = val; r.font.name = FONT
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE
        else:
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if ri%2 else LIGHT
            if ci == 0:
                r.font.size = Pt(17); r.font.bold = True
                r.font.color.rgb = RISK if row[0]=="1" else (EXIGO if row[0]=="5" else NAVY)
            else:
                r.font.size = Pt(12.5); r.font.color.rgb = INK if ci==1 else SLATE
                r.font.bold = (ci==1)

# ---------------------------------------------------------------- SLIDE 6  payments split
image_slide(6, "payments-split.png")
# ---------------------------------------------------------------- SLIDE 7  payments target
image_slide(7, "payments-target.png")

# ---------------------------------------------------------------- SLIDE 8  vendor speed
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "You don't control your own speed", accent=AMBER)
tf = textbox(s, 0.75, 1.95, 11.8, 4.8)
para(tf, [("Your customer and distributor websites are built on an aging “starter-kit” that came with Exigo — meant to bootstrap onto Exigo's own flows, not to be a long-term product platform.", False, INK)], 18, first=True, space_after=16, line=1.25)
for a,b in [
    ("Old software only the vendor fully understands — ", "presentation, business rules and data access are all tangled together, with thin documentation."),
    ("Even small changes are slow and expensive — ", "any change needs vendor know-how and waits in the vendor's queue."),
    ("You don't set your own timeline or cost — ", "your roadmap runs on the vendor's priorities, not yours."),
]:
    para(tf, [("•  ", True, AMBER), (a, True, INK), (b, False, SLATE)], 17, space_after=12, line=1.22)
rect(s, 0.75, 6.0, 11.8, 0.75, fill=RGBColor(0xFB,0xF1,0xDF))
tf2 = textbox(s, 1.0, 6.0, 11.3, 0.75, anchor=MSO_ANCHOR.MIDDLE)
para(tf2, [("This is why you feel stuck “firefighting.”", True, RGBColor(0x9a,0x6b,0x15))], 16, first=True)

# ---------------------------------------------------------------- SLIDE 9  three experiences
image_slide(9, "three-experiences.png")

# ---------------------------------------------------------------- SLIDE 10  missing basics
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "The missing basics")
items = [
    ("No proper access control", "Who can see and do what isn't clearly defined — a security and oversight gap between customer, distributor, admin and support."),
    ("No single sign-on", "More passwords and inconsistent security — both a customer annoyance and a liability."),
    ("Analytics that don't yet drive decisions", "You have dashboards, but not yet a clear answer to “how healthy is the business right now?”"),
]
y = 2.0
for head, body in items:
    rect(s, 0.75, y, 11.85, 1.35, fill=LIGHT)
    rect(s, 0.75, y, 0.12, 1.35, fill=NAVY)
    tf = textbox(s, 1.1, y, 11.3, 1.35, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(head, True, NAVY)], 20, first=True, space_after=3)
    para(tf, [(body, False, SLATE)], 15.5, line=1.2)
    y += 1.55
tf2 = textbox(s, 0.75, 6.75, 11.8, 0.5)
para(tf2, [("None of these caused the architecture gap — but they make every fix slower and riskier.", False, MUTED)], 14, first=True)

# ---------------------------------------------------------------- SLIDE 11  no in-house team
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "The root cause: no in-house technology team", accent=RISK)
tf = textbox(s, 0.75, 2.0, 11.8, 3.0)
para(tf, [("You work through capable ", False, INK), ("vendors", True, INK), (", but have ", False, INK), ("no technical counterpart of your own", True, RISK), (".", False, INK)], 22, first=True, space_after=16, line=1.25)
for a,b in [
    ("No one to own the code ", "or direct the vendors — so change waits in their queue and you stay in reactive mode."),
    ("This is the upstream cause ", "of most of the other issues. With no one on your side, change is outsourced by default."),
    ("Architecture alone won't fix it — ", "better systems help, but can't substitute for having someone on your side to drive them."),
]:
    para(tf, [("•  ", True, RISK), (a, True, INK), (b, False, SLATE)], 18, space_after=12, line=1.22)
rect(s, 0.75, 6.05, 11.8, 0.75, fill=RGBColor(0xFB,0xE9,0xE8))
tf2 = textbox(s, 1.0, 6.05, 11.3, 0.75, anchor=MSO_ANCHOR.MIDDLE)
para(tf2, [("This is why the staffing decision — coming up — gates everything.", True, RISK)], 16, first=True)

# ---------------------------------------------------------------- SLIDES 12–15  images
image_slide(12, "target-architecture.png")
image_slide(13, "source-of-truth.png")
image_slide(14, "roadmap.png")
image_slide(15, "staffing.png")

# ---------------------------------------------------------------- SLIDE 16  what we need
s = slide(); rect(s,0,0,SW,SH,fill=WHITE)
title_block(s, "What we'd need from leadership")
steps = [
    ("1", "Direction on staffing", "Build, partner, or hybrid — this is the gate."),
    ("2", "Go-ahead for a small first step", "Fixed-scope and low-risk: an integration map, early billing/compliance protections, and a sized plan for the rest."),
    ("3", "A few confirmations  [to confirm]", "Current card-handling & compliance setup, how subscriptions split across Exigo/Fluid, what Fluid supports, and the exact Web App boundary."),
]
y = 1.95
for num, head, body in steps:
    rect(s, 0.75, y, 0.85, 1.35, fill=NAVY)
    tfn = textbox(s, 0.75, y, 0.85, 1.35, anchor=MSO_ANCHOR.MIDDLE)
    para(tfn, [(num, True, WHITE)], 30, align=PP_ALIGN.CENTER, first=True)
    rect(s, 1.6, y, 11.0, 1.35, fill=LIGHT)
    tf = textbox(s, 1.9, y, 10.5, 1.35, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(head, True, NAVY)], 20, first=True, space_after=3)
    para(tf, [(body, False, SLATE)], 15.5, line=1.2)
    y += 1.55

# ---------------------------------------------------------------- SLIDE 17  closing
s = slide(); rect(s,0,0,SW,SH,fill=NAVY)
rect(s, 0, 0, 0.28, SH, fill=TEAL)
tf = textbox(s, 1.1, 1.5, 11.2, 4.6)
para(tf, [("The right two platforms. A real business.", True, WHITE)], 30, first=True, space_after=16, line=1.1)
para(tf, [("What's missing is ", False, RGBColor(0xCF,0xD9,0xE6)), ("ownership of the layer in between", True, TEAL), (" — and the ", False, RGBColor(0xCF,0xD9,0xE6)), ("capability to change it.", True, TEAL)], 22, space_after=18, line=1.2)
para(tf, [("The path is low-risk and incremental, the first step is small, and the payoff is substantial: lower risk, faster delivery, a better experience, and the freedom to innovate.", False, RGBColor(0xD7,0xDF,0xE8))], 18, space_after=18, line=1.3)
rect(s, 1.12, 5.75, 11.1, 0.95, fill=NAVY2)
tf2 = textbox(s, 1.4, 5.75, 10.6, 0.95, anchor=MSO_ANCHOR.MIDDLE)
para(tf2, [("The sooner the staffing decision is made, the sooner the firefighting stops.", True, WHITE)], 19, first=True, line=1.15)

# chrome on content slides (2..16; skip title 1 and closing 17)
for i, sl in enumerate(prs.slides, start=1):
    if i in (1, 17):
        continue
    chrome(sl, i)

prs.save(OUT)
print("Saved:", OUT, "| slides:", len(prs.slides))
