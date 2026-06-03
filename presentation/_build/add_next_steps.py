import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PPTX = r'C:\Users\vmank\Documents\Code\GBClients\LimbicArc-assessment\limbic-arc-assessment\presentation\Limbic Arc Tech Assessment -  Briefing.pptx'

NAVY  = RGBColor(0x1F,0x3A,0x5F)
NAVY2 = RGBColor(0x2C,0x52,0x7E)
INK   = RGBColor(0x16,0x20,0x2B)
SLATE = RGBColor(0x5C,0x6B,0x7A)
MUTED = RGBColor(0x8A,0x97,0xA4)
TEAL  = RGBColor(0x1A,0xA1,0x92)
RISK  = RGBColor(0xD9,0x53,0x4F)
AMBER = RGBColor(0xE8,0xA3,0x3D)
LIGHT = RGBColor(0xF3,0xF6,0xF9)
WHITE = RGBColor(0xFF,0xFF,0xFF)
LINE  = RGBColor(0xD7,0xDE,0xE6)
FONT  = "Segoe UI"

prs = Presentation(PPTX)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

def rect(s, l, t, w, h, fill=None, line=None, line_w=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w if line_w else 1)
    return sp

def textbox(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf, m, Pt(2))
    return tf

def run(p, text, bold=False, color=INK, size=16):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = FONT
    return r

def para(tf, segs, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
         space_after=8, space_before=0, line=1.12, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    if isinstance(segs, str):
        segs = [(segs, bold, color)]
    for seg in segs:
        txt = seg[0]; b = seg[1] if len(seg)>1 else bold; c = seg[2] if len(seg)>2 else color
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = c; r.font.name = FONT
    return p

s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, fill=WHITE)

# Title
tf = textbox(s, 0.7, 0.5, 12, 1.0)
para(tf, "Immediate next steps", 31, NAVY, bold=True, first=True, line=1.05)
rect(s, 0.72, 1.42, 1.7, 0.07, fill=TEAL)

# Step cards — each is a numbered block + detail column
steps = [
    (
        "1",
        NAVY,
        "Decide how to staff it",
        "This week",
        "Choose from three options: hire your own team (3–6 months to ramp), "
        "bring in a partner (fastest start), or start with a partner and grow your "
        "own team over time (hybrid). The rest of the plan waits on this decision.",
    ),
    (
        "2",
        TEAL,
        "Kick off a small, fixed-scope first step",
        "Immediately after staffing decision",
        "Deliverables: (a) Document the current Exigo–Fluid integration and data "
        "flows — a map the team can actually work from. "
        "(b) Add audit logging on all profile, subscription and payment writes. "
        "(c) Stop any new card tokenization happening on the Exigo side while the "
        "full fix is scoped. "
        "(d) Produce a properly sized plan and timeline for the phases that follow.",
    ),
    (
        "3",
        AMBER,
        "Confirm a few specifics with your teams & vendors",
        "In parallel",
        "Four items still need owner confirmation before the plan is finalized: "
        "(a) Your current card-handling and PCI compliance setup. "
        "(b) How subscriptions are split between Exigo and Fluid today. "
        "(c) What Fluid supports through its API for hosted card capture and "
        "subscription events. "
        "(d) The exact boundary of the Web App — specifically the single-login "
        "touchpoint.",
    ),
]

y = 1.62
row_h = 1.52
for num, col, head, timing, body in steps:
    # number badge
    rect(s, 0.7, y, 0.72, row_h, fill=col)
    tfn = textbox(s, 0.7, y, 0.72, row_h, anchor=MSO_ANCHOR.MIDDLE)
    para(tfn, num, 28, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)

    # timing tag
    rect(s, 1.42, y, 11.18, row_h, fill=LIGHT)
    tft = textbox(s, 11.6, y+0.14, 1.6, 0.36)
    pt = tft.paragraphs[0]; pt.alignment = PP_ALIGN.RIGHT
    r = pt.add_run(); r.text = timing
    r.font.size = Pt(12); r.font.color.rgb = MUTED; r.font.name = FONT; r.font.italic = True

    # heading + body
    tf2 = textbox(s, 1.65, y, 10.8, row_h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf2, [(head, True, NAVY)], 19, first=True, space_after=4)
    para(tf2, [(body, False, SLATE)], 13.5, line=1.22, space_after=0)

    y += row_h + 0.14

# footer
tf3 = textbox(s, 0.7, 7.06, 9, 0.35)
para(tf3, "Limbic Arc — Platform Assessment · CEO Briefing", 9.5, MUTED, first=True)
tf4 = textbox(s, 11.8, 7.06, 1.0, 0.35)
n = len(prs.slides)
para(tf4, str(n), 9.5, MUTED, align=PP_ALIGN.RIGHT, first=True)

prs.save(PPTX)
print(f"Saved. Total slides: {len(prs.slides)}")
